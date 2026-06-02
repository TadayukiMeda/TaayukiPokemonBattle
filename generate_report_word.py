from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── Color palette ──────────────────────────────────────────────
BLUE_DARK   = RGBColor(0x0F, 0x2D, 0x5E)
BLUE_MID    = RGBColor(0x1A, 0x56, 0xDB)
BLUE_LIGHT  = RGBColor(0x3B, 0x82, 0xF6)
BLUE_PALE   = RGBColor(0x93, 0xC5, 0xFD)
BLUE_TINT   = RGBColor(0xDB, 0xEA, 0xFE)
BLUE_BG     = RGBColor(0xEF, 0xF6, 0xFF)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_MAIN   = RGBColor(0x1E, 0x3A, 0x5F)
TEXT_MUTED  = RGBColor(0x64, 0x74, 0x8B)
RED_SOFT    = RGBColor(0xFE, 0xF2, 0xF2)
RED_DARK    = RGBColor(0x99, 0x1B, 0x1B)
GREEN_DARK  = RGBColor(0x16, 0x65, 0x34)
YELLOW_BG   = RGBColor(0xFF, 0xFB, 0xEB)
YELLOW_DARK = RGBColor(0x78, 0x35, 0x0F)
GRAY_LIGHT  = RGBColor(0xF1, 0xF5, 0xF9)
GRAY_TEXT   = RGBColor(0x47, 0x55, 0x69)

# ── XML helpers ────────────────────────────────────────────────

def set_cell_bg(cell, rgb: RGBColor):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    hex_color = '{:02X}{:02X}{:02X}'.format(rgb[0], rgb[1], rgb[2])
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    existing = tcPr.find(qn('w:shd'))
    if existing is not None:
        tcPr.remove(existing)
    tcPr.append(shd)


def set_cell_border(cell, top=None, bottom=None, left=None, right=None):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if val is None:
            continue
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), val.get('val', 'single'))
        el.set(qn('w:sz'), str(val.get('sz', 4)))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), val.get('color', '000000'))
        existing = tcBorders.find(qn(f'w:{side}'))
        if existing is not None:
            tcBorders.remove(existing)
        tcBorders.append(el)


def set_table_no_spacing(table):
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    spacing = OxmlElement('w:tblCellSpacing')
    spacing.set(qn('w:w'), '0')
    spacing.set(qn('w:type'), 'dxa')
    existing = tblPr.find(qn('w:tblCellSpacing'))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(spacing)


def set_para_spacing(para, before=0, after=0, line=None):
    pPr = para._p.get_or_add_pPr()
    spacing = pPr.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        pPr.append(spacing)
    spacing.set(qn('w:before'), str(before))
    spacing.set(qn('w:after'), str(after))
    if line:
        spacing.set(qn('w:line'), str(line))
        spacing.set(qn('w:lineRule'), 'auto')


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn('w:tcMar'))
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for side, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        existing = tcMar.find(qn(f'w:{side}'))
        if existing is not None:
            tcMar.remove(existing)
        tcMar.append(el)


def add_run(para, text, bold=False, italic=False, size=None, color=None, font_name=None):
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if font_name:
        run.font.name = font_name
    return run


def cell_para(cell, text='', bold=False, size=10, color=None, align=WD_ALIGN_PARAGRAPH.LEFT,
              before=40, after=40):
    if cell.paragraphs and cell.paragraphs[0].text == '':
        para = cell.paragraphs[0]
    else:
        para = cell.add_paragraph()
    para.alignment = align
    set_para_spacing(para, before=before, after=after)
    if text:
        add_run(para, text, bold=bold, size=size, color=color)
    return para


# ── Section heading ────────────────────────────────────────────

def add_section_heading(doc, num, title_ja, title_en):
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_no_spacing(tbl)
    tbl.columns[0].width = Cm(1.4)
    tbl.columns[1].width = Cm(16.6)
    row = tbl.rows[0]

    # number circle (simulated with dark blue cell)
    c0 = row.cells[0]
    set_cell_bg(c0, BLUE_DARK)
    set_cell_margins(c0, 40, 40, 60, 60)
    p = cell_para(c0, str(num), bold=True, size=14, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=60, after=60)

    c1 = row.cells[1]
    set_cell_bg(c1, BLUE_TINT)
    set_cell_margins(c1, 30, 30, 140, 60)
    p2 = cell_para(c1, title_ja, bold=True, size=14, color=BLUE_DARK, before=30, after=0)
    p3 = c1.add_paragraph()
    set_para_spacing(p3, before=0, after=30)
    add_run(p3, title_en, size=8, color=TEXT_MUTED)

    doc.add_paragraph()


def add_sub_heading(doc, text):
    p = doc.add_paragraph()
    set_para_spacing(p, before=120, after=60)
    add_run(p, text, bold=True, size=12, color=BLUE_DARK)
    # underline via border on paragraph bottom
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '{:02X}{:02X}{:02X}'.format(BLUE_TINT[0], BLUE_TINT[1], BLUE_TINT[2]))
    pBdr.append(bottom)
    pPr.append(pBdr)


# ── Highlight box (dark blue) ──────────────────────────────────

def add_highlight_box(doc, heading, items, body_text=None):
    tbl = doc.add_table(rows=1, cols=1)
    set_table_no_spacing(tbl)
    tbl.columns[0].width = Cm(18)
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, BLUE_DARK)
    set_cell_margins(cell, 120, 120, 180, 180)

    p_head = cell_para(cell, heading, bold=True, size=11, color=BLUE_PALE, before=0, after=60)

    if body_text:
        p_body = cell.add_paragraph()
        set_para_spacing(p_body, before=0, after=80)
        add_run(p_body, body_text, size=10, color=WHITE)

    for item in items:
        p_item = cell.add_paragraph()
        set_para_spacing(p_item, before=20, after=20)
        add_run(p_item, '▸  ', size=10, color=BLUE_PALE)
        add_run(p_item, item, size=10, color=WHITE)

    doc.add_paragraph()


# ── Info table (2-col key/value) ───────────────────────────────

def add_info_table(doc, rows_data, col_widths=(4.5, 13.5)):
    tbl = doc.add_table(rows=len(rows_data), cols=2)
    set_table_no_spacing(tbl)
    tbl.columns[0].width = Cm(col_widths[0])
    tbl.columns[1].width = Cm(col_widths[1])
    border_style = {'val': 'single', 'sz': 4, 'color': 'DBEAFE'}

    for i, (key, val) in enumerate(rows_data):
        row = tbl.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        bg = BLUE_BG if i % 2 == 0 else WHITE
        set_cell_bg(c0, BLUE_BG)
        set_cell_bg(c1, bg)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        for cell in (c0, c1):
            set_cell_border(cell, top=border_style, bottom=border_style,
                            left=border_style, right=border_style)
        cell_para(c0, key, bold=True, size=10, color=BLUE_DARK)
        cell_para(c1, val, size=10, color=TEXT_MAIN)

    doc.add_paragraph()


# ── KPI row ────────────────────────────────────────────────────

def add_kpi_row(doc, kpis):
    n = len(kpis)
    w_each = 18.0 / n
    tbl = doc.add_table(rows=1, cols=n)
    set_table_no_spacing(tbl)
    border_style = {'val': 'single', 'sz': 4, 'color': 'DBEAFE'}

    for i, (num, label) in enumerate(kpis):
        col_width = Cm(w_each - 0.2)
        tbl.columns[i].width = col_width
        cell = tbl.rows[0].cells[i]
        set_cell_bg(cell, WHITE)
        set_cell_margins(cell, 100, 100, 80, 80)
        set_cell_border(cell, top={'val': 'single', 'sz': 12, 'color': '1A56DB'},
                        bottom=border_style, left=border_style, right=border_style)
        p_num = cell_para(cell, num, bold=True, size=16, color=BLUE_MID,
                          align=WD_ALIGN_PARAGRAPH.CENTER, before=60, after=20)
        p_label = cell.add_paragraph()
        p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para_spacing(p_label, before=0, after=60)
        add_run(p_label, label, size=9, color=TEXT_MUTED)

    doc.add_paragraph()


# ── Findings grid (2 cols) ─────────────────────────────────────

def add_findings_grid(doc, left_title, left_items, right_title, right_items):
    tbl = doc.add_table(rows=1, cols=2)
    set_table_no_spacing(tbl)
    tbl.columns[0].width = Cm(8.8)
    tbl.columns[1].width = Cm(8.8)
    border_style = {'val': 'single', 'sz': 4, 'color': 'E2E8F0'}

    for col_idx, (title, items, title_color) in enumerate([
        (left_title, left_items, GREEN_DARK),
        (right_title, right_items, RED_DARK)
    ]):
        cell = tbl.rows[0].cells[col_idx]
        set_cell_bg(cell, WHITE)
        set_cell_margins(cell, 120, 120, 160, 160)
        set_cell_border(cell, top=border_style, bottom=border_style,
                        left=border_style, right=border_style)

        p_title = cell_para(cell, title, bold=True, size=11, color=title_color, before=0, after=80)
        for item_text in items:
            p = cell.add_paragraph()
            set_para_spacing(p, before=40, after=40)
            add_run(p, '  ' + item_text, size=10, color=GRAY_TEXT)

    doc.add_paragraph()


# ── Timeline ───────────────────────────────────────────────────

def add_timeline_item(doc, date, title, items, urgent=False):
    p_date = doc.add_paragraph()
    set_para_spacing(p_date, before=100, after=20)
    add_run(p_date, date, bold=True, size=9, color=BLUE_MID)

    p_title = doc.add_paragraph()
    set_para_spacing(p_title, before=0, after=40)
    add_run(p_title, title, bold=True, size=12, color=BLUE_DARK)
    if urgent:
        add_run(p_title, '  【要対応】', bold=True, size=9, color=RGBColor(0xDC, 0x26, 0x26))

    tbl = doc.add_table(rows=1, cols=1)
    set_table_no_spacing(tbl)
    tbl.columns[0].width = Cm(17)
    cell = tbl.rows[0].cells[0]
    set_cell_bg(cell, WHITE)
    set_cell_margins(cell, 80, 80, 160, 120)
    border = {'val': 'single', 'sz': 4, 'color': 'DBEAFE'}
    left_border = {'val': 'single', 'sz': 12, 'color': '1A56DB'}
    set_cell_border(cell, top=border, bottom=border, left=left_border, right=border)

    for i, item in enumerate(items):
        before = 0 if i > 0 else 20
        p = cell_para(cell, '・ ' + item, size=10, color=GRAY_TEXT,
                      before=before, after=20) if i > 0 else \
            cell_para(cell, '・ ' + item, size=10, color=GRAY_TEXT, before=20, after=20)

    doc.add_paragraph()


# ── Company target table ───────────────────────────────────────

def add_target_table(doc, companies):
    headers = ['企業名 / 概要', 'A', 'B', 'C', 'D', 'E', 'F', 'G', '優先度', 'メモ']
    widths = [3.8, 0.55, 0.55, 0.55, 0.55, 0.55, 0.55, 0.55, 2.2, 7.7]

    tbl = doc.add_table(rows=1 + len(companies), cols=len(headers))
    set_table_no_spacing(tbl)
    for i, w in enumerate(widths):
        tbl.columns[i].width = Cm(w)

    border_style = {'val': 'single', 'sz': 4, 'color': 'DBEAFE'}

    # Header row
    hrow = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        set_cell_bg(cell, BLUE_DARK)
        set_cell_margins(cell, 60, 60, 80, 80)
        set_cell_border(cell, top=border_style, bottom=border_style,
                        left=border_style, right=border_style)
        align = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
        cell_para(cell, h, bold=True, size=9, color=BLUE_PALE,
                  align=align, before=40, after=40)

    # Data rows
    for ri, company in enumerate(companies):
        row = tbl.rows[ri + 1]
        row_bg = BLUE_BG if ri % 2 == 0 else WHITE
        name, desc, badges, priority, memo = company

        for ci in range(len(headers)):
            cell = row.cells[ci]
            set_cell_bg(cell, row_bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            set_cell_border(cell, top=border_style, bottom=border_style,
                            left=border_style, right=border_style)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Name + desc
        c0 = row.cells[0]
        p_name = cell_para(c0, name, bold=True, size=10, color=BLUE_DARK, before=40, after=10)
        p_desc = c0.add_paragraph()
        set_para_spacing(p_desc, before=0, after=40)
        add_run(p_desc, desc, size=8, color=TEXT_MUTED)

        # Badges A-G
        BADGE_MAP = {
            'check': ('✓', GREEN_DARK),
            'partial': ('△', RGBColor(0xD9, 0x77, 0x06)),
            'unknown': ('?', TEXT_MUTED),
            'cross': ('✗', RGBColor(0xDC, 0x26, 0x26)),
        }
        for bi, badge_key in enumerate(badges):
            sym, color = BADGE_MAP.get(badge_key, ('?', TEXT_MUTED))
            c = row.cells[bi + 1]
            cell_para(c, sym, bold=(badge_key in ('check', 'cross')), size=11, color=color,
                      align=WD_ALIGN_PARAGRAPH.CENTER, before=40, after=40)

        # Priority
        cell_para(row.cells[8], priority, bold=True, size=9, color=BLUE_MID,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=40, after=40)

        # Memo
        cell_para(row.cells[9], memo, size=9, color=GRAY_TEXT, before=40, after=40)

    doc.add_paragraph()


# ── Competitor table ───────────────────────────────────────────

def add_competitor_table(doc, companies):
    headers = ['企業名', '拠点', '特徴', '分類', 'Sanheとの比較ポイント']
    widths = [4.0, 1.5, 4.5, 2.5, 5.5]

    tbl = doc.add_table(rows=1 + len(companies), cols=len(headers))
    set_table_no_spacing(tbl)
    for i, w in enumerate(widths):
        tbl.columns[i].width = Cm(w)

    border_style = {'val': 'single', 'sz': 4, 'color': 'FECACA'}

    hrow = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        set_cell_bg(cell, RGBColor(0x7F, 0x1D, 0x1D))
        set_cell_margins(cell, 60, 60, 80, 80)
        set_cell_border(cell, top=border_style, bottom=border_style,
                        left=border_style, right=border_style)
        cell_para(cell, h, bold=True, size=9, color=WHITE,
                  align=WD_ALIGN_PARAGRAPH.LEFT, before=40, after=40)

    for ri, (name, desc, hq, feature, cls_text, note) in enumerate(companies):
        row = tbl.rows[ri + 1]
        row_bg = RED_SOFT if ri % 2 == 0 else WHITE
        for ci in range(len(headers)):
            cell = row.cells[ci]
            set_cell_bg(cell, row_bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            set_cell_border(cell, top=border_style, bottom=border_style,
                            left=border_style, right=border_style)

        c0 = row.cells[0]
        cell_para(c0, name, bold=True, size=10, color=RED_DARK, before=40, after=10)
        p_d = c0.add_paragraph()
        set_para_spacing(p_d, before=0, after=40)
        add_run(p_d, desc, size=8, color=TEXT_MUTED)

        cell_para(row.cells[1], hq, size=10, color=TEXT_MAIN, before=40, after=40)
        cell_para(row.cells[2], feature, size=9, color=GRAY_TEXT, before=40, after=40)
        cell_para(row.cells[3], cls_text, bold=True, size=9, color=RED_DARK, before=40, after=40)
        cell_para(row.cells[4], note, size=9, color=GRAY_TEXT, before=40, after=40)

    doc.add_paragraph()


# ── Roadmap table ──────────────────────────────────────────────

def add_roadmap_table(doc, rows_data):
    headers = ['展示会', '時期', '規模', '精度', '優先アクション']
    widths = [4.5, 2.0, 2.0, 1.5, 8.0]

    tbl = doc.add_table(rows=1 + len(rows_data), cols=len(headers))
    set_table_no_spacing(tbl)
    for i, w in enumerate(widths):
        tbl.columns[i].width = Cm(w)

    border_style = {'val': 'single', 'sz': 4, 'color': 'DBEAFE'}

    hrow = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        set_cell_bg(cell, BLUE_DARK)
        set_cell_margins(cell, 60, 60, 80, 80)
        set_cell_border(cell, top=border_style, bottom=border_style,
                        left=border_style, right=border_style)
        cell_para(cell, h, bold=True, size=10, color=BLUE_PALE, before=40, after=40)

    for ri, row_data in enumerate(rows_data):
        row = tbl.rows[ri + 1]
        bg = BLUE_BG if ri % 2 == 0 else WHITE
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            set_cell_bg(cell, bg)
            set_cell_margins(cell, 50, 50, 80, 80)
            set_cell_border(cell, top=border_style, bottom=border_style,
                            left=border_style, right=border_style)
            bold = ci == 0
            cell_para(cell, val, bold=bold, size=10, color=TEXT_MAIN, before=40, after=40)

    doc.add_paragraph()


# ══════════════════════════════════════════════════════════════
#  BUILD DOCUMENT
# ══════════════════════════════════════════════════════════════

doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# Default paragraph style
style = doc.styles['Normal']
style.font.name = 'メイリオ'
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.space_before = Pt(0)


# ── COVER ──────────────────────────────────────────────────────

cover_tbl = doc.add_table(rows=1, cols=1)
set_table_no_spacing(cover_tbl)
cover_tbl.columns[0].width = Cm(18)
cover_cell = cover_tbl.rows[0].cells[0]
set_cell_bg(cover_cell, BLUE_DARK)
set_cell_margins(cover_cell, 400, 400, 400, 400)

# Badge line
p_badge = cell_para(cover_cell, 'CONFIDENTIAL  —  INTERNAL PITCH DOCUMENT',
                    size=9, color=BLUE_PALE, before=0, after=200)
p_badge.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Show name
p_show = cell_para(cover_cell, 'ライフスタイルWeek夏 2026', bold=True,
                   size=22, color=BLUE_PALE, before=0, after=40)

# Main title
p_title = cell_para(cover_cell, 'Sanhe 展示会アプローチ戦略', bold=True,
                    size=28, color=WHITE, before=0, after=80)

# Subtitle
p_sub = cell_para(cover_cell, 'Japan Market Entry — Trade Show Intelligence Report',
                  size=11, color=RGBColor(0xBF, 0xDB, 0xFE), before=0, after=300)

# KPI strip inside cover
kpi_items = [
    ('6 / 24', '開催開始日'),
    ('1,250社', '総出展社数'),
    ('42,000名', '年間来場者数'),
    ('71%', '購買決定権保有率'),
]
kpi_tbl = OxmlElement('w:tbl')
# We'll do it simply as a line of text pairs
for num, label in kpi_items:
    p_kpi = cover_cell.add_paragraph()
    p_kpi.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_spacing(p_kpi, before=20, after=20)
    add_run(p_kpi, f'  {num}  ', bold=True, size=14, color=BLUE_PALE)
    add_run(p_kpi, f'{label}    ', size=9, color=RGBColor(0xBF, 0xDB, 0xFE))

# Footer line
p_footer = cover_cell.add_paragraph()
p_footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_para_spacing(p_footer, before=300, after=0)
add_run(p_footer, 'Sanhe Interfashion — Japan B2B Strategy    ',
        size=9, color=RGBColor(0x93, 0xC5, 0xFD))
add_run(p_footer, '作成日：2026年5月26日    ',
        size=9, color=RGBColor(0x93, 0xC5, 0xFD))
add_run(p_footer, 'For Internal Use Only',
        size=9, color=RGBColor(0x93, 0xC5, 0xFD))

doc.add_page_break()


# ── PAGE 1: 展示会概要 ─────────────────────────────────────────

add_section_heading(doc, 1, '展示会概要', 'SHOW OVERVIEW')

add_info_table(doc, [
    ('展示会名', 'ライフスタイルWeek 夏 2026（第29回）'),
    ('主催', 'RX Japan 合同会社'),
    ('会期', '2026年6月24日（水）〜 26日（金）　3日間'),
    ('会場', '東京ビッグサイト（東展示棟）'),
    ('出展規模', '関連10展 合計 1,250社・38,000点'),
    ('来場者数', '毎年 約42,000名（71.2%が購買決定権保有バイヤー）'),
    ('性格', 'B2B 商談専門展（一般来場不可）'),
    ('来場登録', '無料 / 事前登録締切：2026年6月19日（金）'),
    ('国際来場者', '中国・韓国・台湾・タイ・ベトナム・米国・フランス等'),
])

add_sub_heading(doc, '構成展示会（10展）')

expo_data = [
    ('★ Sanhe重点ターゲット', '国際 ファッション雑貨 EXPO',
     'バッグ・アクセサリー・革小物・帽子などファッション雑貨専門。OEM商談に最適。'),
    ('★ 要確認', 'インバウンド向けグッズ EXPO',
     '外国人観光客向けの日本土産・工芸品系。本革×和柄ブランドが出展する場合あり。'),
    ('参考調査', '国際 雑貨 EXPO',
     'セレクトショップ・雑貨店向け。生活雑貨・インテリア小物など幅広い出展。'),
    ('参考調査', '推し活 EXPO',
     'グッズ・キャラクターコラボ系。革小物活用の可能性を探る。'),
    ('対象外', '国際 インテリア EXPO',
     '家具・インテリア雑貨専門。Sanheのターゲット外。'),
    ('対象外', 'コスメ・食器・鍋 EXPO 他',
     '美容・キッチン系。直接的な接点なし。'),
]

expo_tbl = doc.add_table(rows=len(expo_data), cols=3)
set_table_no_spacing(expo_tbl)
expo_tbl.columns[0].width = Cm(2.8)
expo_tbl.columns[1].width = Cm(5.0)
expo_tbl.columns[2].width = Cm(10.2)
border_style = {'val': 'single', 'sz': 4, 'color': 'DBEAFE'}

for ri, (tag, name, desc) in enumerate(expo_data):
    row = expo_tbl.rows[ri]
    is_target = '★' in tag
    bg = BLUE_BG if is_target else (GRAY_LIGHT if '参考' in tag else WHITE)
    tag_color = BLUE_MID if is_target else (TEXT_MUTED if '対象外' in tag else RGBColor(0x5B, 0x8F, 0xD6))

    for ci in range(3):
        cell = row.cells[ci]
        set_cell_bg(cell, bg)
        set_cell_margins(cell, 60, 60, 100, 100)
        set_cell_border(cell, top=border_style, bottom=border_style,
                        left=border_style, right=border_style)

    cell_para(row.cells[0], tag, bold=True, size=8, color=tag_color, before=40, after=40)
    cell_para(row.cells[1], name, bold=True, size=10, color=TEXT_MAIN, before=40, after=40)
    cell_para(row.cells[2], desc, size=9, color=GRAY_TEXT, before=40, after=40)

doc.add_paragraph()
doc.add_page_break()


# ── PAGE 2: 戦略的意義 ────────────────────────────────────────

add_section_heading(doc, 2, 'Sanheにとっての戦略的意義', 'WHY THIS SHOW MATTERS')

add_highlight_box(doc,
    '▌ 本展示会がSanheの日本市場開拓に直結する理由',
    [
        '購買決定権保有バイヤー比率 71.2% — 担当者ではなくオーナーに直接話せる',
        '競合OEM工場も出展しており、Sanheの差別化を「対比」で見せられる',
        'セレクトショップ・バッグ専門店のバイヤーから「どこのOEM使ってるか」を逆聞きできる',
        '国際バイヤーも来場 → 日本以外のアジア市場への横展開情報も収集可能',
    ],
    body_text='Sanheのターゲット（日本の革バッグ・財布ブランド企業）のバイヤーおよびオーナーが、年42,000名規模で来場する場。メールでは届かない「意思決定者への直接アプローチ」が3日間で実現できる機会。'
)

add_kpi_row(doc, [
    ('3日間', '商談集中期間'),
    ('71%', '決定権保有バイヤー率'),
    ('10社+/日', 'アプローチ目標'),
    ('6/19', '無料登録締切'),
])

add_findings_grid(doc,
    '✅ この展示会でできること',
    [
        'OEM外注先を探している日本ブランドに直接ピッチ',
        '競合OEM工場（中国・台湾・ベトナム）の価格・品質を調査',
        'バイヤーとの名刺交換 → 後日メールフォローの許可を取得',
        '「今の工場に何か不満はあるか」を聞く市場調査',
        '海外展開中の日本ブランドを現場で特定',
    ],
    '⚠️ 現状の課題・限界',
    [
        '出展者ディレクトリは事前登録なしにはアクセス困難（ログイン必須）',
        'ファッション雑貨セクションはOEM競合が多く、ターゲットブランドが少なかった',
        'バッグ・財布専門のセクション絞り込みが事前必須',
        '来場登録締切まで残り約3週間（6月19日）',
    ]
)

doc.add_page_break()


# ── PAGE 3: 出展者分析（10社）────────────────────────────────

add_section_heading(doc, 3, '出展者分析（調査済み10社）',
                    'EXHIBITOR ANALYSIS — 10 COMPANIES IDENTIFIED')

# info note
note_tbl = doc.add_table(rows=1, cols=1)
set_table_no_spacing(note_tbl)
note_tbl.columns[0].width = Cm(18)
nc = note_tbl.rows[0].cells[0]
set_cell_bg(nc, YELLOW_BG)
set_cell_margins(nc, 80, 80, 160, 120)
set_cell_border(nc, top={'val': 'single', 'sz': 4, 'color': 'FDE68A'},
                bottom={'val': 'single', 'sz': 4, 'color': 'FDE68A'},
                left={'val': 'single', 'sz': 12, 'color': 'F59E0B'},
                right={'val': 'single', 'sz': 4, 'color': 'FDE68A'})
cell_para(nc, 'ℹ  調査範囲について', bold=True, size=10, color=YELLOW_DARK, before=0, after=40)
note_text = nc.add_paragraph()
set_para_spacing(note_text, before=0, after=0)
add_run(note_text,
        '「ファッション雑貨EXPO」「インバウンド向けグッズEXPO」セクションの出展者データ（全3バッチ）を分析。'
        'ログイン認証のためシステムによる全件取得は不可。全1,250社中の抽出分析。',
        size=9, color=YELLOW_DARK)
doc.add_paragraph()

# ── Unified 10-company table ──────────────────────────────────
# Columns: 企業名/所在地 | 主要製品 | OEM/ODM | 製造拠点 | 分類 | 商談のポイント | 優先度
COL_WIDTHS = [3.5, 2.4, 1.2, 2.2, 2.2, 5.5, 1.0]
HEADERS    = ['企業名 / 所在地', '主要製品', 'OEM/\nODM', '製造拠点', '分類', '商談のポイント', '優先度']
BD = {'val': 'single', 'sz': 4, 'color': 'DBEAFE'}
BD_RED = {'val': 'single', 'sz': 4, 'color': 'FECACA'}

# section divider helper
def add_divider_row(tbl, label, bg_hex, text_hex):
    row = tbl.add_row()
    merged = row.cells[0]
    for i in range(1, len(COL_WIDTHS)):
        merged = merged.merge(row.cells[i])
    set_cell_bg(merged, RGBColor(int(bg_hex[0:2],16), int(bg_hex[2:4],16), int(bg_hex[4:6],16)))
    set_cell_margins(merged, 60, 60, 140, 80)
    tc = merged._tc
    tcPr = tc.get_or_add_tcPr()
    cell_para(merged, '▌  ' + label, bold=True, size=9,
              color=RGBColor(int(text_hex[0:2],16), int(text_hex[2:4],16), int(text_hex[4:6],16)),
              before=40, after=40)

def add_data_row(tbl, name, loc, product, oem, base, cls_label, cls_color, note, priority_text, bg_color):
    row = tbl.add_row()
    for ci, w in enumerate(COL_WIDTHS):
        tbl.columns[ci].width = Cm(w)
    cells = row.cells

    # col 0: name + location
    set_cell_bg(cells[0], bg_color)
    set_cell_margins(cells[0], 60, 60, 100, 80)
    set_cell_border(cells[0], top=BD, bottom=BD, left=BD, right=BD)
    p_name = cell_para(cells[0], name, bold=True, size=10, color=BLUE_DARK, before=40, after=10)
    p_loc  = cells[0].add_paragraph()
    set_para_spacing(p_loc, before=0, after=40)
    add_run(p_loc, loc, size=8, color=TEXT_MUTED)

    # cols 1-5
    data = [product, oem, base, cls_label, note]
    for ci, val in enumerate(data, start=1):
        set_cell_bg(cells[ci], bg_color)
        set_cell_margins(cells[ci], 60, 60, 80, 80)
        set_cell_border(cells[ci], top=BD, bottom=BD, left=BD, right=BD)
        color = cls_color if ci == 4 else GRAY_TEXT
        bold  = ci == 4
        align = WD_ALIGN_PARAGRAPH.CENTER if ci in (2, 3) else WD_ALIGN_PARAGRAPH.LEFT
        cell_para(cells[ci], val, bold=bold, size=9, color=color, align=align, before=40, after=40)

    # col 6: priority
    set_cell_bg(cells[6], bg_color)
    set_cell_margins(cells[6], 60, 60, 60, 60)
    set_cell_border(cells[6], top=BD, bottom=BD, left=BD, right=BD)
    cell_para(cells[6], priority_text, bold=True, size=10, color=BLUE_MID,
              align=WD_ALIGN_PARAGRAPH.CENTER, before=40, after=40)

# Build table (header + 10 data rows + 3 section dividers)
tbl10 = doc.add_table(rows=1, cols=len(COL_WIDTHS))
set_table_no_spacing(tbl10)
for ci, w in enumerate(COL_WIDTHS):
    tbl10.columns[ci].width = Cm(w)

# Header
hrow = tbl10.rows[0]
for ci, h in enumerate(HEADERS):
    c = hrow.cells[ci]
    set_cell_bg(c, BLUE_DARK)
    set_cell_margins(c, 60, 60, 80, 80)
    set_cell_border(c, top=BD, bottom=BD, left=BD, right=BD)
    cell_para(c, h, bold=True, size=9, color=BLUE_PALE,
              align=WD_ALIGN_PARAGRAPH.CENTER, before=40, after=40)

# Section 1: Targets
add_divider_row(tbl10, 'アプローチターゲット（商談候補）', '0f2d5e', '93c5fd')

targets = [
    ('株式会社 ラビット', '東京・中野区 / 創業48年',
     'バッグ・財布・革小物', 'あり', '中国（協力工場\n社員常駐）',
     '★★★ 最優先', BLUE_MID,
     '自社ブランド「Mineed by Osel」保有。中国工場依存 → 韓国切替提案の余地大。経営者に直接ピッチ可。',
     '🔵🔵🔵', BLUE_BG),
    ('株式会社 HKM', '東京・浅草橋 / SIXWHEELSLIFE',
     '本革バッグ・財布・ポーチ', 'ODM可', '外注（詳細未確認）',
     '★★ 優先', BLUE_MID,
     '京都西陣織×本革の差別化ブランド。ODM対応可・インバウンドEXPOに出展 → 海外展開志向あり。',
     '🔵🔵', BLUE_BG),
    ('株式会社 HARU', '所在地調査中',
     'Pig Skin革製品\nバッグ・ポーチ', 'ODM可', 'Made in Japan主軸',
     '★★ 優先', BLUE_MID,
     '代理店募集中・ODM対応可。外注先拡大に前向きな可能性。Made in Japanからの分散提案が刺さりうる。',
     '🔵🔵', WHITE),
    ('株式会社 シフレ', '所在地調査中',
     'スーツケース・バッグ全般', 'あり', '中国工場',
     '★ 参考', TEXT_MUTED,
     'OEM/ODM事業部あり。スーツケース主力のため革小物の割合要確認。中国→韓国の切替メリットを訴求。',
     '🔵', BLUE_BG),
    ('株式会社 エヌ・エックス', '所在地調査中 / Ryugasaki4716',
     '牛革ベルト・革小物', '要確認', '国内製造訴求',
     '★ 参考', TEXT_MUTED,
     '自社ブランド保有。国内製造訴求が強いためOEM移行ハードル高め。小物特化で量産規模が小さい可能性。',
     '🔵', WHITE),
    ("D'JORA Lifestyles", 'シンガポール法人 / 日本出展',
     'バッグ・レザー製品', '要確認', '外注（詳細未確認）',
     '★ 参考', TEXT_MUTED,
     '英語商談可能。東京拠点なし → 優先度低め。ただし海外展開済みのため韓国工場との相性を探る価値あり。',
     '🔵', BLUE_BG),
]
for t in targets:
    add_data_row(tbl10, *t)

# Section 2: Partner candidates
add_divider_row(tbl10, '代理店・パートナー候補（日本側OEM仲介）', '1e40af', 'bfdbfe')

partners = [
    ('株式会社 茅', '日本 / OEM仲介業',
     'バッグ・小物OEM仲介', 'あり', '中国（自社工場）',
     '代理店候補', RGBColor(0x1E, 0x40, 0xAF),
     '中国工場を自社保有するOEM仲介。競合だが日本側代理店として取り込める可能性あり。中国→韓国への切替提案から関係構築。',
     '🟡', RGBColor(0xF0, 0xF9, 0xFF)),
    ('株式会社 ユーティ', '日本 / OEM受託',
     'バッグ・雑貨OEM受託', 'あり', '中国工場ネットワーク',
     '代理店候補', RGBColor(0x1E, 0x40, 0xAF),
     '日本語対応OEM仲介。中国工場依存 → Sanheを供給元として提案。日本ブランドへの間接アプローチルートになりうる。',
     '🟡', WHITE),
]
for p in partners:
    add_data_row(tbl10, *p)

# Section 3: Competitor research
add_divider_row(tbl10, '競合調査対象（価格・品質ベンチマーク用）', '7f1d1d', 'fecaca')

competitors = [
    ('HAINING ALLBRIGHT', '中国・浙江省 / OEM工場',
     'バッグ・財布OEM', 'あり', '中国',
     '直接競合', RED_DARK,
     '低価格・大ロット対応。価格帯・MOQ・リードタイムを調査してSanheの差別化ポジションを数値化する。',
     '🔴', RED_SOFT),
    ('RIDER ENTERPRISE CO.', '台湾 / 台中・中国・ベトナム',
     'バッグ・小物OEM', 'あり', '台湾・中国・ベトナム',
     '直接競合', RED_DARK,
     '3カ国工場体制が強み。多拠点対応 vs Sanheの「韓国品質特化」の差別化軸を整理する比較材料。',
     '🔴', WHITE),
]
for c in competitors:
    add_data_row(tbl10, *c)

doc.add_paragraph()

# Legend row
p_leg = doc.add_paragraph()
set_para_spacing(p_leg, before=40, after=80)
add_run(p_leg, '優先度：', bold=True, size=9, color=TEXT_MAIN)
add_run(p_leg, ' 🔵🔵🔵 最優先商談  ', size=9, color=TEXT_MUTED)
add_run(p_leg, '🔵🔵 優先訪問  ', size=9, color=TEXT_MUTED)
add_run(p_leg, '🔵 可能なら接触  ', size=9, color=TEXT_MUTED)
add_run(p_leg, '🟡 関係構築候補  ', size=9, color=TEXT_MUTED)
add_run(p_leg, '🔴 競合調査のみ', size=9, color=TEXT_MUTED)

add_highlight_box(doc,
    '▌ 分析サマリー',
    [
        '最優先：株式会社ラビット（創業48年・中国工場依存からの切替余地が最も大きい）',
        '競合2社は商談ではなく「ブース立ち寄り・カタログ収集・価格確認」のみで十分',
        '全1,250社のうち未調査セクションが多数残存 → 来場後のフロア探索でさらに候補を増やせる',
    ],
    body_text='調査した出展者（全3バッチ）の中から10社を選定。うちターゲット商談候補6社、代理店候補2社、競合調査対象2社。'
)

doc.add_page_break()


# ── PAGE 4: アクションプラン ──────────────────────────────────

add_section_heading(doc, 4, '当日アクションプラン', 'ON-SITE ACTION PLAN')

add_sub_heading(doc, '当日の動き方 — 4ステップフロー')

steps = [
    ('Step 1', '競合マッピング', '中国・台湾OEM工場ブースを先に回る。価格・MOQ・リードタイム・品質ポジションを把握。'),
    ('Step 2', 'ターゲット特定', 'バッグ・革小物系の日本ブランドブースを探す。ペルソナ条件（年商・東京・海外展開）を目視確認。'),
    ('Step 3', 'ファーストコンタクト', '「Sanheのカタログ＋サンプル」を持参。30秒ピッチ後、名刺交換と次回アポを打診。'),
    ('Step 4', '展示会後フォロー', '72時間以内にメール送信。韓国工場見学の招待状を添付。温度感の高い企業を優先。'),
]

steps_tbl = doc.add_table(rows=1, cols=4)
set_table_no_spacing(steps_tbl)
for i in range(4):
    steps_tbl.columns[i].width = Cm(4.3)
border_style = {'val': 'single', 'sz': 4, 'color': 'DBEAFE'}

for i, (step_num, step_name, step_desc) in enumerate(steps):
    cell = steps_tbl.rows[0].cells[i]
    set_cell_bg(cell, BLUE_BG)
    set_cell_margins(cell, 100, 100, 120, 120)
    set_cell_border(cell, top={'val': 'single', 'sz': 12, 'color': '1A56DB'},
                    bottom=border_style, left=border_style, right=border_style)
    cell_para(cell, step_num, bold=True, size=9, color=BLUE_MID,
              align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=20)
    cell_para(cell, step_name, bold=True, size=11, color=BLUE_DARK,
              align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=40)
    p_desc = cell.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_para_spacing(p_desc, before=0, after=60)
    add_run(p_desc, step_desc, size=9, color=GRAY_TEXT)

doc.add_paragraph()

add_sub_heading(doc, '事前準備チェックリスト')

add_timeline_item(doc, '今週中 — 5月30日まで', '来場者登録', [
    'lifestyle-expo.jp より来場者登録を完了（無料）',
    '登録時に「バイヤー」「仕入れ・OEM商談目的」を選択',
    '締切：2026年6月19日（金）— 過ぎると有料になる可能性あり',
], urgent=True)

add_timeline_item(doc, '6月1日〜10日', '出展者リスト事前調査', [
    '来場者登録後、ディレクトリにアクセスし「バッグ」「革小物」「財布」カテゴリで絞り込み',
    'ペルソナ条件A〜Gに近い出展者を10〜15社リストアップ',
    '各社の公式サイトで年商・所在地・海外展開状況を事前確認',
])

add_timeline_item(doc, '6月10日〜20日', 'ピッチ資料・持参物の準備', [
    'Sanhe会社概要（日本語版）— A4両面、1枚',
    '革サンプル（財布・ハンドバッグ用素材スウォッチ）を持参',
    '価格表（MOQ・リードタイム・対応素材）日本語版',
    '名刺（日本語表記、韓国工場の写真入り推奨）',
])

add_timeline_item(doc, '6月24日〜26日（当日）', '展示会訪問・商談実施', [
    '国際ファッション雑貨EXPOを重点的に回る（東展示棟）',
    '1日10社アプローチ目標（3日間で30社）',
    '名刺にメモ書き → その日の夜に整理',
    '競合OEM工場のカタログ・価格情報を収集',
])

add_timeline_item(doc, '6月27日〜30日（展示会後）', 'フォローアップ送信', [
    '名刺交換した全社にお礼メール送信（72時間以内）',
    '温度感の高い企業に韓国工場見学の案内を送付',
    '次の展示会（FaW TOKYO 10月）に向けたアポ打診',
])

doc.add_page_break()


# ── PAGE 5: ロードマップ ──────────────────────────────────────

add_section_heading(doc, 5, '展示会ロードマップ 2026', 'TRADE SHOW ROADMAP')

roadmap_rows = [
    ('ライフスタイルWeek夏', '6/24〜26', '1,250社', '★★☆', '来場登録（6/19締切）— 最優先'),
    ('FaW TOKYO 秋', '10/7〜9', '850社', '★★★', '7月に出展者リスト公開予定 → チェック'),
    ('ハンドバッグかばん見本市（TBL）', '11/10〜', 'バッグ専門', '★★★', '8月に来場登録・事前アポ打診'),
    ('TOKYO LEATHER FAIR', '12/3〜4', '55社', '★★☆', '素材工場・職人系 → 競合情報収集'),
    ('START EXHIBITION（浅草）', '未定', '23社', '★★★', '東日本バッグ工業組合に問い合わせ'),
]
add_roadmap_table(doc, roadmap_rows)

add_highlight_box(doc,
    '▌ 推奨戦略 — 3ステップアプローチ',
    [
        'Step 1（6月）ライフスタイルWeek夏 — 競合調査・日本市場の肌感をつかむ。名刺30枚収集目標。',
        'Step 2（10月）FaW TOKYO秋 — 6月で得た情報をもとにピッチを洗練。「具体的に提案できる」状態で臨む。',
        'Step 3（11月）ハンドバッグかばん見本市 — ペルソナ直撃。本命商談。韓国工場見学アポを複数社から取る。',
    ]
)

# Footer
p_foot = doc.add_paragraph()
p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para_spacing(p_foot, before=200, after=0)
add_run(p_foot, 'Sanhe Interfashion — Japan Market Entry Strategy', size=9, color=TEXT_MUTED)
p_foot2 = doc.add_paragraph()
p_foot2.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_para_spacing(p_foot2, before=20, after=0)
add_run(p_foot2, '本資料は社内ピッチ用途に限定。無断転載・外部共有禁止。作成日：2026年5月26日', size=8, color=RGBColor(0xAA, 0xBB, 0xCC))


doc.save('ライフスタイルWeek報告書_Sanhe.docx')
print('Done: ライフスタイルWeek報告書_Sanhe.docx')
